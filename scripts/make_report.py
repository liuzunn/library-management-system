# -*- coding: utf-8 -*-
"""生成《实验一 建立相关数据库》实验报告 DOCX。

以 实验报告（上机）模板.docx 为底板：保留封面与表格样式，
把表格主体重建为六个部分并填入完整内容。
"""
import copy
import io
import json
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import _Row

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_helpers import (HEI, KAI, MONO, SONG, clean_row, clear_cell,
                            code_block, data_table, figure, heading, para,
                            set_run_font)

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(r"C:\Users\LZ\Desktop\软工")
FIG = ROOT / "figures"
SQL = ROOT / "sql"

stats = json.loads((ROOT / "_work" / "stats.json").read_text(encoding="utf-8"))
verify = json.loads((ROOT / "_work" / "verify.json").read_text(encoding="utf-8"))
vmap = {v["title"]: v for v in verify}

TPL = ROOT / "实验报告（上机）模板.docx"
OUT = ROOT / "实验报告（上机）-图书管理系统数据库设计.docx"

doc = docx.Document(str(TPL))

# ---------------------------------------------------------------- 封面
ps = doc.paragraphs
ps[5].runs[1].text = "软件工程"
ps[7].runs[1].text = "实验一  建立相关数据库"
for idx in (9, 11, 13):
    ps[idx].runs[1].text = " " * 26
ps[16].runs[0].text = " 2026 年 9 月 23 日"

# ---------------------------------------------------------------- 重建表格
tbl = doc.tables[0]
template_tr = copy.deepcopy(tbl.rows[0]._tr)        # 整宽内容行
sig_tr = copy.deepcopy(tbl.rows[-1]._tr)            # 教师签字行
for r in list(tbl.rows):
    r._tr.getparent().remove(r._tr)


def new_row():
    tr = copy.deepcopy(template_tr)
    tbl._tbl.append(tr)
    row = _Row(tr, tbl)
    clean_row(row)
    cell = row.cells[0]
    clear_cell(cell)
    return cell


def body(cell, text, **kw):
    kw.setdefault("first_line", 0.21)
    kw.setdefault("space_after", 3)
    return para(cell, text, **kw)


def bullets(cell, items, start=1, size=10.5):
    for i, t in enumerate(items, start):
        body(cell, "%d. %s" % (i, t), size=size)


# =====================================================================
# 一、实验目的及要求
# =====================================================================
c = new_row()
para(c, "一、实验目的及要求", size=13, bold=True, name=HEI, space_before=0, space_after=6)
para(c, "（一）实验目的", size=11, bold=True, name=HEI, space_after=3)
bullets(c, [
    "掌握实体-联系图（E-R 图）的设计方法，能够正确区分实体、属性与联系，并标注联系的种类；",
    "掌握数据字典的设计方法，能够用数据项、数据结构、数据存储和处理过程来描述系统的数据组成；",
    "能够根据 E-R 图和数据字典设计出系统的数据库逻辑结构，并用关系数据库管理系统加以实现；",
    "掌握数据流图的绘制方法，能够自顶向下、逐层分解地描述系统的数据加工过程；",
    "掌握在数据库管理系统中建立数据库、基本表以及实现数据录入、修改、删除等维护操作的方法。",
])
para(c, "（二）实验要求", size=11, bold=True, name=HEI, space_before=6, space_after=3)
body(c, "以“图书管理系统”为题目，完成下列三项工作：")
bullets(c, [
    "设计图书管理系统的实体-联系图，并用数据字典描述其中的数据元素（自行编制）；",
    "在数据库管理平台上建立图书管理系统的数据库，建立“图书”“读者”“借阅”三张基本表，并完成图书信息与读者信息的录入、删除和修改功能；",
    "设计并绘制图书管理系统的数据流图（自行设计）。",
])
para(c, "（三）关于开发平台的说明", size=11, bold=True, name=HEI, space_before=6, space_after=3)
body(c, "实验指导书给出的开发平台为 Microsoft Access 或 Visual Basic 6.0。本次实验在 "
        "Windows 11 环境下改用 MySQL 8.0.44 数据库服务器实现，客户端使用 Navicat "
        "Premium 17。改用 MySQL 的原因见“五、结果分析”部分。所设计的 E-R 图、数据字典"
        "和数据流图与具体数据库产品无关，因此概念设计与逻辑设计的结论完全一致。",
     align="justify")

# =====================================================================
# 二、软、硬件环境
# =====================================================================
c = new_row()
para(c, "二、软、硬件环境", size=13, bold=True, name=HEI, space_before=0, space_after=6)
para(c, "（一）硬件环境", size=11, bold=True, name=HEI, space_after=3)
data_table(c, ["项目", "配置"], [
    ["处理器", "AMD Ryzen 9 7945HX（16 核 32 线程）"],
    ["内存", "16 GB（系统可用 15.7 GB）"],
    ["硬盘", "Samsung MZVL21T0HCLR-00BL2 固态硬盘，954 GB"],
    ["显示分辨率", "1920 × 1080 及以上"],
], widths=[1.5, 4.8])
para(c, "（二）软件环境", size=11, bold=True, name=HEI, space_before=8, space_after=3)
data_table(c, ["类别", "软件名称", "版本"], [
    ["操作系统", "Microsoft Windows 11 家庭版（64 位）", "10.0.26200"],
    ["数据库服务器", "MySQL Community Server（winx64）", "8.0.44"],
    ["数据库客户端", "Navicat Premium 17", "17.0.4"],
    ["数据库客户端", "MySQL Workbench 8.0 CE", "8.0.44"],
    ["绘图与统计", "Python + matplotlib", "3.13.9 / 3.10.6"],
    ["数据库驱动", "PyMySQL（Python 访问 MySQL）", "2.2.8"],
], widths=[1.35, 3.35, 1.60])
body(c, "数据库字符集统一采用 utf8mb4，排序规则采用 utf8mb4_unicode_ci，以保证中文数据的"
        "正确存储与比较；服务器监听端口为 3306。",
     space_before=6, align="justify")

# =====================================================================
# 三、算法或原理分析（实验内容）
# =====================================================================
c = new_row()
para(c, "三、算法或原理分析（实验内容）", size=13, bold=True, name=HEI, space_before=0, space_after=6)

heading(c, "3.1  需求分析")
body(c, "图书管理系统面向图书管理员和读者两类用户，需要管理图书信息、读者信息以及二者之间的"
        "借阅关系。经分析，系统的功能需求如下：")
data_table(c, ["功能模块", "功能描述"], [
    ["图书信息管理", "图书信息的录入、修改、删除与查询"],
    ["读者信息管理", "读者信息的录入、修改、删除与查询"],
    ["借阅管理", "办理借书与还书，登记借书时间和还书时间"],
    ["查询与统计", "按出版社、读者类型、时间等条件统计图书借阅情况"],
], widths=[1.7, 4.6])
body(c, "系统的数据需求及完整性要求如下：", space_before=6)
bullets(c, [
    "图书信息包括：书号、书名、作者、价格、出版社、摘要；",
    "读者信息包括：卡号、姓名、性别、单位、类型、级别；",
    "借阅信息包括：书号、卡号、借书时间、还书时间；",
    "每本图书的书号唯一且非空，每位读者的卡号唯一且非空；",
    "一条借阅记录必须对应一个已存在的书号和一个已存在的卡号（参照完整性）；",
    "还书时间不得早于借书时间；同一本图书在未归还期间不得重复借出。",
])

heading(c, "3.2  概念结构设计——实体-联系图（E-R 图）")
body(c, "根据需求分析，系统包含“图书”和“读者”两个实体。图书的属性为书号、书名、作者、"
        "价格、出版社和摘要，其中书号为标识属性；读者的属性为卡号、姓名、性别、单位、类型"
        "和级别，其中卡号为标识属性。图书与读者之间通过“借阅”发生联系，该联系为 1:N 联系"
        "（一位读者可以借阅多本图书，一本图书在某一时刻至多由一位读者持有），联系的属性为"
        "借书时间和还书时间。所设计的 E-R 图如图 1 所示。", align="justify")
figure(c, FIG / "fig1_er_diagram.png", width_in=5.35, caption="图 1  图书管理系统 E-R 图")

heading(c, "3.3  数据字典")
body(c, "数据字典是对数据流图中各元素定义的集合，包括数据项、数据结构、数据存储和处理过程"
        "四类条目。根据上述需求，自行编制数据字典如下。", align="justify")
para(c, "表 1  数据项（数据元素）定义", size=9.5, bold=True, name=HEI, space_before=6, space_after=2)
data_table(c, ["编号", "数据项名称", "别名", "类型与长度", "取值范围"], [
    ["I01", "书号", "book_no", "VARCHAR(20)", "B + 4 位数字"],
    ["I02", "书名", "title", "VARCHAR(120)", "非空字符串"],
    ["I03", "作者", "author", "VARCHAR(60)", "非空字符串"],
    ["I04", "价格", "price", "DECIMAL(8,2)", "0.00 ~ 999999.99"],
    ["I05", "出版社", "publisher", "VARCHAR(80)", "字符串"],
    ["I06", "摘要", "summary", "VARCHAR(500)", "字符串"],
    ["I07", "卡号", "card_no", "VARCHAR(20)", "R + 4 位数字"],
    ["I08", "姓名", "name", "VARCHAR(30)", "非空字符串"],
    ["I09", "性别", "gender", "ENUM(‘男’,‘女’)", "男 / 女"],
    ["I10", "单位", "unit", "VARCHAR(80)", "字符串"],
    ["I11", "类型", "type", "VARCHAR(20)", "学生 / 教师 / 职工"],
    ["I12", "级别", "grade", "VARCHAR(20)", "本科 / 硕士 / 博士 / 教职工"],
    ["I13", "借书时间", "borrow_date", "DATE", "yyyy-mm-dd"],
    ["I14", "还书时间", "return_date", "DATE", "yyyy-mm-dd 或 NULL"],
], widths=[0.55, 1.25, 1.25, 1.45, 1.80], size=8.5)

para(c, "表 2  数据结构定义", size=9.5, bold=True, name=HEI, space_before=8, space_after=2)
data_table(c, ["编号", "数据结构名称", "组成"], [
    ["DS01", "图书信息", "书号 + 书名 + 作者 + 价格 + 出版社 + 摘要"],
    ["DS02", "读者信息", "卡号 + 姓名 + 性别 + 单位 + 类型 + 级别"],
    ["DS03", "借阅信息", "书号 + 卡号 + 借书时间 + 还书时间"],
], widths=[0.7, 1.4, 4.2])

para(c, "表 3  数据存储定义", size=9.5, bold=True, name=HEI, space_before=8, space_after=2)
data_table(c, ["编号", "名称", "输入数据", "输出数据", "组成"], [
    ["D1", "图书表", "图书信息", "图书信息", "DS01"],
    ["D2", "读者表", "读者信息", "读者信息", "DS02"],
    ["D3", "借阅表", "借阅信息", "借阅信息", "DS03"],
], widths=[0.6, 1.1, 1.4, 1.4, 1.8])

para(c, "表 4  处理过程定义", size=9.5, bold=True, name=HEI, space_before=8, space_after=2)
data_table(c, ["编号", "处理过程", "输入", "输出", "处理逻辑"], [
    ["P1", "图书信息管理", "图书信息", "图书表记录", "对图书表进行插入、修改、删除与查询"],
    ["P2", "读者信息管理", "读者信息", "读者表记录", "对读者表进行插入、修改、删除与查询"],
    ["P3", "借阅管理", "借还书申请", "借阅记录", "登记借还书，校验图书与读者的合法性"],
    ["P4", "查询与统计", "查询条件", "统计报表", "读取三个数据存储并汇总输出"],
], widths=[0.6, 1.3, 1.15, 1.15, 2.1])

heading(c, "3.4  数据流图设计")
body(c, "数据流图从数据传递和加工的角度，以图形方式刻画系统内部的数据流动情况。按照自顶向下、"
        "逐层分解的原则，先画出顶层数据流图，再将其展开为第 1 层数据流图。", align="justify")
body(c, "顶层数据流图把整个图书管理系统看作一个加工（编号 0），系统外部有两个实体：图书管理员"
        "和读者。管理员向系统提供图书信息、读者信息和借还书办理请求，系统向管理员输出处理结果"
        "和统计报表；读者向系统提出借书、还书和查询请求，系统向读者返回借阅结果和查询结果。"
        "顶层数据流图如图 2 所示。", align="justify")
figure(c, FIG / "fig2_dfd_top.png", width_in=6.30, caption="图 2  图书管理系统顶层数据流图")
body(c, "把顶层图中的加工“图书管理系统”继续分解，得到 4 个加工和 3 个数据存储：加工 1 负责"
        "图书信息管理，加工 2 负责读者信息管理，加工 3 负责借阅管理，加工 4 负责查询与统计；"
        "数据存储 D1、D2、D3 分别对应数据库中的图书表、读者表和借阅表。第 1 层数据流图如图 3 "
        "所示。", align="justify")
figure(c, FIG / "fig3_dfd_level1.png", width_in=6.35, caption="图 3  图书管理系统第 1 层数据流图")

heading(c, "3.5  逻辑结构设计与规范化分析")
body(c, "将 E-R 图按“一个实体型转换为一个关系模式、一个联系转换为一个关系模式”的规则转换，"
        "得到如下三个关系模式：", align="justify")
for t in ["图书（书号，书名，作者，价格，出版社，摘要）",
          "读者（卡号，姓名，性别，单位，类型，级别）",
          "借阅（书号，卡号，借书时间，还书时间）"]:
    body(c, t, indent=0.45, first_line=0.0, size=10.5)
body(c, "其中，图书关系的码为书号，读者关系的码为卡号；借阅关系的码为（书号，卡号，借书时间）"
        "的属性组合，书号和卡号分别是参照图书关系和读者关系的外码。三个关系模式在数据库中的"
        "实现及外码联系如图 4 所示。", space_before=4, align="justify")
figure(c, FIG / "fig4_schema.png", width_in=6.35,
       caption="图 4  图书管理系统数据库表结构与外码联系")
para(c, "（1）范式分析", size=10.5, bold=True, name=HEI, space_before=6, space_after=2)
body(c, "第一范式（1NF）：三个关系模式的每个属性都是不可再分的基本数据项，均满足 1NF。", first_line=0.21)
body(c, "第二范式（2NF）：图书关系和读者关系的主码均为单个属性，不存在部分函数依赖；借阅关系"
        "的主码为（书号，卡号，借书时间），非主属性只有还书时间，它完全函数依赖于整个主码，"
        "不存在部分函数依赖。因此三个关系均满足 2NF。")
body(c, "第三范式（3NF）：三个关系中都不存在非主属性对主码的传递函数依赖（例如出版社名称直接"
        "存放在图书关系中，没有通过出版社编号间接决定）。因此三个关系均满足 3NF，设计是合理的。")
para(c, "（2）关于联系种类的说明", size=10.5, bold=True, name=HEI, space_before=6, space_after=2)
body(c, "实验指导书附录二的 E-R 图中，“图书”一侧标注 N、“读者”一侧标注 1，其语义是“一本图书在"
        "同一时刻至多由一位读者借阅，一位读者可以借阅多本图书“。考虑到图书馆业务需要保留历史"
        "借阅记录（同一本书在不同时期可由不同读者借阅），本实验在实现时把借阅关系的主码取为"
        "（书号，卡号，借书时间）的联合主码：这样既能够表达上述 N:1 的“当前借阅”语义，又能够"
        "完整保存 M:N 的历史借阅事实，不会丢失数据。", align="justify")

# =====================================================================
# 四、程序代码或实现过程
# =====================================================================
sql_all = (SQL / "01_create_database.sql").read_text(encoding="utf-8")
parts = re.split(r"/\* -{10,}\n(.*?)\n\s*-{10,} \*/", sql_all, flags=re.S)
blocks = []
for i in range(1, len(parts), 2):
    t = parts[i].strip().split("\n")[0].strip()
    blocks.append((t, parts[i + 1].strip()))
m = re.search(r"DROP DATABASE.*", parts[0], flags=re.S)
create_db = m.group(0).strip()

c = new_row()
para(c, "四、程序代码或实现过程", size=13, bold=True, name=HEI, space_before=0, space_after=6)

heading(c, "4.1  建立数据库")
body(c, "在 MySQL 8.0.44 中建立数据库 library_db，字符集取 utf8mb4，排序规则取 "
        "utf8mb4_unicode_ci，以便正确存储中文数据。", align="justify")
code_block(c, create_db)

heading(c, "4.2  建立基本表")
body(c, "按照 3.5 节得到的关系模式建立三张基本表。建表时把主键、非空、默认值、外码等完整性"
        "约束一并定义在表上，由数据库服务器负责检查，避免把完整性检查放在应用程序中。",
     align="justify")
for title, code in blocks[:3]:
    para(c, title, size=10, bold=True, name=HEI, space_before=6, space_after=2)
    code_block(c, code)

heading(c, "4.3  建立视图")
body(c, "为方便日常查询和统计报表输出，建立了三个视图：借阅明细视图通过三表等值连接把书号、"
        "书名、卡号、姓名等信息组合在一起；图书借阅统计视图按书汇总借阅次数与在借数量；读者"
        "借阅统计视图按读者类型和级别汇总借阅情况。", align="justify")
for title, code in blocks[3:6]:
    para(c, title, size=10, bold=True, name=HEI, space_before=6, space_after=2)
    code_block(c, code)

heading(c, "4.4  建立触发器")
body(c, "借阅表上的触发器在插入借阅记录之前进行校验：如果同一本图书尚未归还，则拒绝重复借出；"
        "如果还书时间早于借书时间，同样拒绝写入。触发器在数据库层保证了业务的正确性。",
     align="justify")
for title, code in blocks[6:7]:
    para(c, title, size=10, bold=True, name=HEI, space_before=6, space_after=2)
    code_block(c, code)

heading(c, "4.5  建立存储过程")
body(c, "实验要求系统完成图书和读者信息管理功能，包括信息的录入、删除和修改。本实验把每一项"
        "操作封装成一个存储过程，客户端只需调用过程名并传入参数即可完成操作，既简化了调用，"
        "又可以在过程内部统一加入业务校验。", align="justify")
for title, code in blocks[7:]:
    para(c, title, size=10, bold=True, name=HEI, space_before=6, space_after=2)
    code_block(c, code)

heading(c, "4.6  初始化实验数据")
body(c, "为了便于验证功能并进行统计分析，编写脚本向三张表载入初始数据：图书 24 种、读者 18 人、"
        "借阅记录 128 条，数据日期范围为 2026 年 3 月至 2026 年 9 月。下面以图书表为例给出插入"
        "语句的写法（完整脚本见随报告提交的 02_sample_data.sql）。", align="justify")
code_block(c, """INSERT INTO book (book_no, title, author, price, publisher, summary) VALUES
  ('B0001', '数据库系统概论', '王珊', 45.00, '高等教育出版社',
   '系统介绍关系数据库原理、SQL语言与数据库设计方法。'),
  ('B0002', '软件工程导论', '张海藩', 39.50, '清华大学出版社',
   '阐述软件生命周期、结构化分析与面向对象方法。'),
  ('B0003', '数据结构(C语言版)', '严蔚敏', 42.00, '清华大学出版社',
   '线性表、树、图及其经典算法的C语言实现。');

INSERT INTO reader (card_no, name, gender, unit, type, grade) VALUES
  ('R0001', '张伟', '男', '计算机科学与技术学院', '学生', '本科'),
  ('R0002', '李娜', '女', '计算机科学与技术学院', '学生', '硕士'),
  ('R0012', '冯军', '男', '计算机科学与技术学院', '教师', '教职工');

INSERT INTO borrow (book_no, card_no, borrow_date, return_date) VALUES
  ('B0007', 'R0001', '2026-03-06', '2026-04-14'),
  ('B0006', 'R0003', '2026-09-03', '2026-09-12');""")

heading(c, "4.7  在 Navicat Premium 17 中使用数据库")
bullets(c, [
    "启动 Navicat Premium 17，单击“连接”→“MySQL”，新建一个 MySQL 连接；",
    "在连接属性中填写：主机 127.0.0.1、端口 3306、用户名 root、密码 123456，"
    "字符集选择 utf8mb4，然后单击“测试连接”，提示连接成功后保存；",
    "展开连接即可看到 library_db 数据库，其中 Tables 节点下列出 book、reader、borrow "
    "三张表，Views 节点下列出三个视图，Functions 节点下列出全部存储过程；",
    "右键单击某张表选择“设计表”，可以查看和修改字段、主键、外键与索引；选择“打开表”"
    "即可浏览和编辑数据；",
    "单击工具栏的“查询”→“新建查询”，在 SQL 编辑器中输入查询语句，按 Ctrl+R 执行，"
    "结果以表格形式显示在下方。",
])

# =====================================================================
# 五、结果分析
# =====================================================================
def find_key(prefix):
    for k in vmap:
        if k.startswith(prefix):
            return k
    raise KeyError(prefix)


def short_err(prefix, limit=70):
    e = vmap[find_key(prefix)]["error"] or ""
    e = re.sub(r"^(IntegrityError|DataError|OperationalError|pymysql\.err\.[A-Za-z]+): ", "", e)
    return e if len(e) <= limit else e[:limit] + " …"


def vtable(cell, prefix, caption=None, size=8.5, max_rows=99):
    v = vmap[find_key(prefix)]
    if caption:
        para(cell, caption, size=9.5, bold=True, name=HEI, space_before=6, space_after=2)
    rows = v["rows"][:max_rows]
    data_table(cell, v["cols"], rows, size=size)
    return v


c = new_row()
para(c, "五、结果分析", size=13, bold=True, name=HEI, space_before=0, space_after=6)

heading(c, "5.1  数据库对象与表结构验证")
body(c, "建库脚本执行完毕后，用 SHOW TABLES 查看数据库中已建立的对象，可以看到 3 张基本表和 "
        "3 个视图；再用 DESCRIBE 查看各表的字段定义，结果与 3.5 节设计的关系模式完全一致。",
     align="justify")
vtable(c, "1. ", caption="表 5  数据库中已建立的对象")
vtable(c, "2. ", caption="表 6  图书表 book 的字段定义")
vtable(c, "3. ", caption="表 7  读者表 reader 的字段定义")
vtable(c, "4. ", caption="表 8  借阅表 borrow 的字段定义")
vtable(c, "6. ", caption="表 9  借阅表上的外键约束")

heading(c, "5.2  图书信息管理功能测试")
body(c, "按照实验要求，对图书信息的录入、修改、删除三项功能逐一进行测试。测试通过调用 4.5 节"
        "建立的存储过程完成，每一步都用查询语句确认执行结果。", align="justify")
vtable(c, "8. ", caption="表 10  录入功能测试：调用 sp_book_insert 插入新书 B0025 后的查询结果")
vtable(c, "10. ", caption="表 11  修改功能测试：调用 sp_book_update 后的查询结果")
vtable(c, "17. ", caption="表 12  删除功能测试：调用 sp_book_delete 后的查询结果")
body(c, "测试结果表明：三种操作均正确执行。录入后可以查到新书 B0025 的完整信息；修改后书名、"
        "价格等字段被正确更新；删除后该书号在图书表中的记录数为 0，且借阅表中与之相关的记录"
        "因外键的 ON DELETE RESTRICT 规则而不会产生孤立记录。读者信息的录入、修改、删除由 "
        "sp_reader_insert、sp_reader_update、sp_reader_delete 三个过程实现，测试过程与图书"
        "信息完全相同，此处不再重复。", align="justify")
vtable(c, "5. ", caption="表 13  借阅明细视图 v_borrow_detail 的查询结果（前 6 行）")

heading(c, "5.3  完整性约束与触发器测试")
body(c, "为了检验约束是否真正生效，特意输入了 5 组非法数据。测试结果如表 14 所示，所有非法"
        "操作都被数据库服务器拒绝，并给出了明确的中文提示。", align="justify")
para(c, "表 14  完整性约束与触发器测试结果", size=9.5, bold=True, name=HEI, space_before=6, space_after=2)
data_table(c, ["序号", "测试内容", "预期结果", "实际返回信息"], [
    ["1", "插入已存在的书号 B0001", "违反实体完整性", short_err("11.")],
    ["2", "借阅不存在的图书 B9999", "违反参照完整性", short_err("12.", 96)],
    ["3", "读者性别填写“未知”", "违反用户定义完整性", short_err("13.")],
    ["4", "把尚未归还的图书再次借出", "被触发器拒绝", short_err("14.")],
    ["5", "还书时间早于借书时间", "被触发器拒绝", short_err("15.")],
], widths=[0.5, 1.85, 1.35, 2.6], size=8.5)
body(c, "由表 14 可见，主键约束、外键约束和 ENUM 取值约束都由数据库服务器自动实施，触发器则"
        "把“图书未归还不得重复借出”“还书时间不得早于借书时间”这两条业务规则固化在数据库层。"
        "这样即使更换了前端程序，规则依然有效，这是把完整性约束交给数据库管理的重要理由。",
     space_before=4, align="justify")

heading(c, "5.4  借阅数据统计分析")
body(c, "在完成功能验证后，利用库中的 128 条借阅记录进行统计分析。统计查询通过前面建立的三个"
        "视图和聚合函数完成，绘图由 Python 的 matplotlib 库实现，统计结果如图 5 所示。",
     align="justify")
figure(c, FIG / "fig5_stats.png", width_in=6.30,
       caption="图 5  图书借阅数据统计分析结果")
vtable(c, "18. ", caption="表 15  数据总量统计")
vtable(c, "20. ", caption="表 16  按读者类型和级别的借阅统计（v_reader_stat 视图）")
vtable(c, "19. ", caption="表 17  图书借阅排行（前 8 名，v_book_stat 视图）")
body(c, "对统计结果的分析如下：", space_before=4)
bullets(c, [
    "学生是图书借阅的主体。18 位读者中有 13 位是学生，共借阅 100 次，占全部借阅量的 78.1%；"
    "教师 4 人借阅 25 次，占 19.5%；职工 3 次，占 2.3%。",
    "借阅周期与读者类型密切相关。学生平均借阅 28.8 天，教师平均 65.0 天，职工平均 75.0 天。"
    "学生借阅频次高、周期短，与课程按章节推进的学习节奏一致；教师和职工借阅周期长，多为"
    "备课、科研的长期占用。",
    "借阅量存在明显的学期周期性。3 月开学季借阅量最高（30 次），7 月暑假期间回落到 13 次，"
    "8 月下旬随新学期临近又回升到 19 次。9 月仅统计到 9 月 20 日，8 次的结果并不代表全月"
    "水平，说明统计报表需要注明数据的截止时间。",
    "从图书来源看，高等教育出版社的图书被借阅 32 次，占 25.0%，居第一位，这与馆藏中教材类"
    "图书占比较高有关；其后依次是机械工业出版社（17 次）和清华大学出版社（14 次）。",
    "借阅次数最多的图书是《数据结构(C语言版)》（8 次），其次是《离散数学》和《大学英语"
    "精读(第三册)》（各 7 次），均为课程配套教材。",
    "截至统计时点仍有 24 条借阅记录尚未归还，其中学生 13 条、教师 10 条、职工 1 条。"
    "结合平均借阅天数可以看出，教师和职工中存在较多的长期占用现象，建议系统补充“应还日期”"
    "和“超期提醒”功能。",
])

heading(c, "5.5  关于开发平台选择的说明")
body(c, "实验指导书建议使用的平台是 Microsoft Access 或 Visual Basic 6.0。Access 属于桌面"
        "数据库，需要依赖 Microsoft Office 环境，在并发访问、用户权限管理和数据量较大时的"
        "性能方面能力有限；VB 6.0 发布于 1998 年，微软已于 2008 年停止对其提供支持，在 "
        "Windows 11 上安装和运行都存在兼容性问题。", align="justify")
body(c, "本实验改用 MySQL 8.0.44 作为数据库服务器、Navicat Premium 17 作为图形化客户端。"
        "MySQL 完整支持主键、外键、ENUM、触发器、存储过程和视图等特性，可以更充分地覆盖本"
        "实验要求掌握的知识点；Navicat 的建库、建表、数据维护操作方式与 Access 类似，易于"
        "上手。由于 E-R 图、数据字典和数据流图都是与具体数据库产品无关的设计文档，因此更换"
        "平台并不影响本实验概念设计与逻辑设计的结论，反而更贴近目前企业中的实际开发环境。",
     align="justify")

heading(c, "5.6  实验结论与心得体会")
bullets(c, [
    "本实验完整地经历了“需求分析→概念结构设计→逻辑结构设计→物理实现→验证”的数据库设计"
    "过程。E-R 图解决“系统中有哪些数据、它们之间有什么关系”的问题，数据字典进一步把每个"
    "数据元素定义清楚，二者是逻辑设计的直接依据。",
    "把完整性约束定义在数据库层而不是应用程序层，是本次实验最深的体会。主键、外键、ENUM "
    "和触发器在测试中都成功拦截了非法数据；如果把这些检查全部写在应用程序里，一旦更换前端"
    "或有人直接操作数据库，数据的一致性就无法保证。",
    "关于联系的种类，指导书附录二给出的 E-R 图中图书一侧为 N、读者一侧为 1。实现时把借阅"
    "关系的主码取为（书号，卡号，借书时间）的联合主码，既表达了“当前借阅”的 N:1 语义，"
    "又保留了同一本书在不同时期被不同读者借阅的历史记录，说明概念设计中的联系种类需要结合"
    "业务的实际语义来确定。",
    "本次设计仍有不足之处：图书表没有设置“类别/学科”字段，因此只能按出版社而不便按学科"
    "分类统计；图书表没有“馆藏册数”字段，无法描述同一种书拥有多个副本的情况；借阅表缺少"
    "“应还日期”“续借次数”等字段，不能支持超期提醒和罚金计算；此外尚未实现用户账号与"
    "权限管理。这些都可以作为后续实验的改进方向。",
])

# ---------------- 教师签字行 ----------------
tr = copy.deepcopy(sig_tr)
tbl._tbl.append(tr)
row = _Row(tr, tbl)
for tc in row._tr.findall(qn("w:tc")):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is not None:
        for tag in ("w:vMerge", "w:noWrap"):
            el = tcPr.find(qn(tag))
            if el is not None:
                tcPr.remove(el)
cells = row.cells
clear_cell(cells[0])
p = cells[0].add_paragraph()
p.alignment = 1
set_run_font(p.add_run("教师签字"), size=11, bold=True)
clear_cell(cells[2])
p = cells[2].add_paragraph()
p.alignment = 1
set_run_font(p.add_run("年    月    日"), size=11)

doc.save(str(OUT))
print("saved", OUT)
print("sections:", len(tbl.rows), "rows")
